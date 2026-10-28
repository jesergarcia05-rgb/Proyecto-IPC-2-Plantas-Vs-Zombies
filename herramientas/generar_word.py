#!/usr/bin/env python3
"""Exporta el mismo contenido de los manuales PDF a Word editable.
Requiere python-docx, ReportLab y Pillow solo para regenerar documentos.
No convierte cada pagina en una imagen: conserva parrafos, negritas y tablas.
"""
from pathlib import Path
import contextlib
import io
import runpy
import reportlab.platypus as platypus
from reportlab.platypus import Paragraph, Image, Table, PageBreak, Spacer, Preformatted, KeepTogether
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

RAIZ = Path(__file__).resolve().parents[1]


def configurar(doc, titulo):
    seccion = doc.sections[0]
    seccion.page_width = Pt(595.2756)
    seccion.page_height = Pt(841.8898)
    seccion.top_margin = Pt(42)
    seccion.bottom_margin = Pt(56)
    seccion.left_margin = seccion.right_margin = Pt(44)
    seccion.footer_distance = Pt(23)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = Pt(14)
    normal.paragraph_format.space_after = Pt(9)
    idioma = OxmlElement('w:lang')
    idioma.set(qn('w:val'), 'es-GT')
    normal.element.get_or_add_rPr().append(idioma)
    for nombre, tamano in [('Title', 28), ('Heading 1', 16), ('Heading 2', 11.5)]:
        estilo = doc.styles[nombre]
        estilo.font.name = 'Arial'
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor.from_string('205742')
        estilo.paragraph_format.keep_with_next = True
    pie = doc.styles['Caption']
    pie.font.name = 'Arial'
    pie.font.size = Pt(8)
    pie.font.italic = False
    pie.font.color.rgb = RGBColor.from_string('50605B')
    pie.paragraph_format.line_spacing = Pt(11)
    pie.paragraph_format.space_after = Pt(10)
    doc.core_properties.title = titulo
    doc.core_properties.author = 'Proyecto Zombis vs. Plantas'
    doc.core_properties.subject = 'Manual editable de IPC1, Proyecto 2, segundo semestre 2026'
    footer = seccion.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Pt(507.27), WD_TAB_ALIGNMENT.RIGHT)
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run('Zombis vs. Plantas · Proyecto 2 · IPC1 · 2S2026\t')
    r.font.name = 'Arial'; r.font.size = Pt(8)
    campo = OxmlElement('w:fldSimple'); campo.set(qn('w:instr'), 'PAGE')
    footer._p.append(campo)


def parrafo(destino, elemento, celda=False):
    nombre = elemento.style.name
    estilo = {'Titulo': 'Title', 'Seccion': 'Heading 1', 'Subtitulo': 'Heading 2', 'Pie': 'Caption'}.get(nombre, 'Normal')
    p = destino.add_paragraph(style=estilo)
    p.paragraph_format.space_before = Pt(getattr(elemento.style, 'spaceBefore', 0))
    p.paragraph_format.space_after = Pt(getattr(elemento.style, 'spaceAfter', 0))
    p.paragraph_format.line_spacing = Pt(getattr(elemento.style, 'leading', 14))
    for fragmento in elemento.frags:
        if getattr(fragmento, 'lineBreak', False):
            p.add_run().add_break(); continue
        texto = getattr(fragmento, 'text', '')
        if not texto: continue
        r = p.add_run(texto)
        r.bold = bool(getattr(fragmento, 'bold', False))
        # Las fuentes del PDF comparten alias para cursiva; el contenido
        # original no usa cursivas, por lo que no se propaga ese alias.
        r.italic = False
        r.font.name = 'Arial'
        r.font.size = Pt(getattr(fragmento, 'fontSize', 10))
        color = getattr(fragmento, 'textColor', None)
        if color is not None:
            r.font.color.rgb = RGBColor(round(color.red * 255), round(color.green * 255), round(color.blue * 255))
    if celda:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = False
    return p


def exportar(doc, elementos):
    salto_pendiente = False
    for e in elementos:
        if isinstance(e, PageBreak):
            salto_pendiente = True
        elif isinstance(e, Paragraph):
            p = parrafo(doc, e)
            if salto_pendiente:
                p.paragraph_format.page_break_before = True
                salto_pendiente = False
        elif isinstance(e, Image):
            p = doc.add_paragraph()
            if salto_pendiente:
                p.paragraph_format.page_break_before = True
                salto_pendiente = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(e.filename), width=Pt(e.drawWidth), height=Pt(e.drawHeight))
        elif isinstance(e, Table):
            t = doc.add_table(rows=0, cols=e._ncols)
            t.style = 'Table Grid'
            t.autofit = False
            anchos = e._argW
            if any(w is None for w in anchos): anchos = [507.27 / e._ncols] * e._ncols
            for col, ancho in zip(t.columns, anchos): col.width = Pt(ancho)
            for indice, valores in enumerate(e._cellvalues):
                fila = t.add_row()
                no_dividir = OxmlElement('w:cantSplit')
                fila._tr.get_or_add_trPr().append(no_dividir)
                if indice == 0:
                    repetir = OxmlElement('w:tblHeader')
                    fila._tr.get_or_add_trPr().append(repetir)
                for i, valor in enumerate(valores):
                    celda = fila.cells[i]
                    celda.width = Pt(anchos[i])
                    if isinstance(valor, Paragraph):
                        parrafo(celda, valor, True)
                        celda._tc.remove(celda.paragraphs[0]._p)
                    else:
                        celda.paragraphs[0].text = str(valor)
                    if indice == 0:
                        sombra = OxmlElement('w:shd'); sombra.set(qn('w:fill'), 'DCECE4')
                        celda._tc.get_or_add_tcPr().append(sombra)
        elif isinstance(e, Preformatted):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = Pt(11)
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run('\n'.join(e.lines))
            r.font.name = 'Consolas'; r.font.size = Pt(8)
        elif isinstance(e, Spacer):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(max(1, e.height))
            p.add_run().font.size = Pt(1)
        elif isinstance(e, KeepTogether):
            exportar(doc, e._content)
        else:
            raise TypeError('Elemento no soportado: ' + type(e).__name__)


class DocumentoWord:
    def __init__(self, nombre, **opciones):
        self.nombre = Path(nombre).with_suffix('.docx')
        self.opciones = opciones

    def build(self, contenido, **opciones):
        doc = Document()
        configurar(doc, self.opciones.get('title', self.nombre.stem))
        exportar(doc, contenido)
        doc.save(self.nombre)


def main():
    original = platypus.SimpleDocTemplate
    try:
        platypus.SimpleDocTemplate = DocumentoWord
        with contextlib.redirect_stdout(io.StringIO()):
            runpy.run_path(str(RAIZ / 'herramientas/generar_manuales.py'), run_name='__main__')
    finally:
        platypus.SimpleDocTemplate = original
    for nombre in ['Manual_de_usuario.docx', 'Manual_tecnico.docx']:
        print('Word generado:', RAIZ / 'documentacion' / nombre)


if __name__ == '__main__':
    main()
